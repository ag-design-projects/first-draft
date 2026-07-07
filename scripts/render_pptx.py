#!/usr/bin/env python
"""Render an editable PPTX from a first-draft deck spec.

When a template is supplied, this renderer opens that template, removes its demo
slides, preserves its slide masters/layouts/theme, and adds generated content
slides. The preserved layouts remain available in PowerPoint Home > Layout.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_LIGHT = ROOT / "sources" / "brands" / "editorial-research" / "templates" / "editorial-slides-template-light.pptx"
DEFAULT_DARK = ROOT / "sources" / "brands" / "editorial-research" / "templates" / "editorial-slides-template-dark.pptx"


ROLE_LAYOUT_HINTS = {
    "cover": ["title", "cover", "front"],
    "section": ["section", "divider", "chapter"],
    "agenda": ["agenda", "contents"],
    "quote": ["quote"],
    "chart": ["chart", "graph", "data"],
    "roadmap": ["timeline", "roadmap"],
    "closing": ["closing", "end", "thank"],
    "image": ["image", "picture", "photo"],
}


def color(value: str, fallback: str = "#111111") -> RGBColor:
    raw = (value or fallback).strip().lstrip("#")
    if len(raw) != 6:
        raw = fallback.lstrip("#")
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def clean_text(value: object) -> str:
    text = str(value or "")
    return text.replace("\u2014", "-").replace("\u2013", "-")


def remove_template_slides(prs: Presentation) -> None:
    slide_ids = list(prs.slides._sldIdLst)  # noqa: SLF001 - python-pptx has no public API.
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)  # noqa: SLF001


def resolve_template(spec: dict, explicit: str | None) -> Path | None:
    candidates = [
        explicit,
        spec.get("template", {}).get("path") if isinstance(spec.get("template"), dict) else None,
        spec.get("brand", {}).get("template"),
        spec.get("brand", {}).get("templates", {}).get("default") if isinstance(spec.get("brand", {}).get("templates"), dict) else None,
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return Path(candidate)
    theme = (spec.get("theme") or spec.get("brand", {}).get("theme") or "light").lower()
    default = DEFAULT_DARK if theme == "dark" else DEFAULT_LIGHT
    return default if default.exists() else None


def layout_name(layout) -> str:
    return (layout.name or "").lower()


def is_blank_layout(layout) -> bool:
    name = layout_name(layout)
    return "blank" in name or (len(layout.shapes) == 0 and len(layout.placeholders) == 0)


def nonblank_layouts(prs: Presentation):
    return [layout for layout in prs.slide_layouts if not is_blank_layout(layout)]


def first_usable_layout(prs: Presentation):
    layouts = nonblank_layouts(prs)
    if layouts:
        return layouts[0]
    return prs.slide_layouts[0]


def choose_layout(prs: Presentation, slide: dict):
    requested = (slide.get("layout_id") or "").lower()
    usable_layouts = nonblank_layouts(prs)
    if requested:
        for index, layout in enumerate(prs.slide_layouts, start=1):
            if requested in {layout.name.lower(), f"slidelayout{index}".lower()} and not is_blank_layout(layout):
                return layout
    role = (slide.get("role") or "insight").lower()
    hints = ROLE_LAYOUT_HINTS.get(role, [])
    for hint in hints:
        for layout in usable_layouts:
            if hint in layout_name(layout):
                return layout
    return first_usable_layout(prs)


def add_textbox(slide, left, top, width, height, text, font_name, font_size, color_value, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = clean_text(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = color_value
    return box


def add_bullets(slide, left, top, width, height, items, font_name, font_size, color_value):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items[:6]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = clean_text(item)
        paragraph.level = 0
        paragraph.font.name = font_name
        paragraph.font.size = Pt(font_size)
        paragraph.font.italic = False
        paragraph.font.color.rgb = color_value
    return box


def add_visual_placeholder(slide, left, top, width, height, visual, font_name, accent, text_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(245, 245, 245)
    shape.line.color.rgb = accent
    frame = shape.text_frame
    frame.clear()
    label = clean_text(f"{visual.get('type', 'visual')}: {visual.get('description', '')}")
    paragraph = frame.paragraphs[0]
    paragraph.text = label
    paragraph.font.name = font_name
    paragraph.font.size = Pt(18)
    paragraph.font.italic = False
    paragraph.font.color.rgb = text_color


def add_footer(slide, evidence, font_name, text_color):
    if not evidence:
        return
    source_text = "; ".join(
        clean_text(f"{ev.get('source', '')} {ev.get('locator', '')}: {ev.get('claim', '')}")
        for ev in evidence[:3]
    )
    add_textbox(slide, Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.28), source_text, font_name, 7, text_color)


def populate_slide(slide, slide_spec: dict, spec: dict, slide_number: int) -> None:
    brand = spec.get("brand", {})
    colors = brand.get("colors", {})
    fonts = brand.get("fonts", {})
    heading_font = fonts.get("heading", "Helvetica Neue")
    body_font = fonts.get("body", "Helvetica Neue")
    primary = color(colors.get("primary", "#111111"))
    accent = color(colors.get("accent", "#0A66CC"))
    text_color = color(colors.get("text", "#111111"))

    role = (slide_spec.get("role") or "insight").lower()
    title_size = 42 if role == "cover" else 30 if role == "section" else 25
    add_textbox(
        slide,
        Inches(0.55),
        Inches(0.35),
        Inches(11.8),
        Inches(0.25),
        f"{brand.get('name', '')} / {slide_number:02d}",
        body_font,
        8,
        accent,
        bold=True,
    )
    add_textbox(slide, Inches(0.55), Inches(0.75), Inches(11.5), Inches(1.0), slide_spec.get("title", ""), heading_font, title_size, primary, bold=True)

    body = slide_spec.get("body", [])
    visual = slide_spec.get("visual", {}) or {}
    image = slide_spec.get("image", {}) or {}
    wants_visual = visual.get("type") and visual.get("type") != "none"
    wants_image = image and image.get("description") and not str(image.get("description")).startswith("none")

    if role in {"cover", "section", "closing"}:
        add_bullets(slide, Inches(0.75), Inches(2.2), Inches(10.8), Inches(2.0), body, body_font, 20, text_color)
    elif wants_visual or wants_image:
        add_bullets(slide, Inches(0.65), Inches(2.0), Inches(5.35), Inches(3.9), body, body_font, 17, text_color)
        placeholder = visual if wants_visual else {"type": "image", "description": image.get("description")}
        add_visual_placeholder(slide, Inches(6.45), Inches(2.0), Inches(5.45), Inches(3.9), placeholder, body_font, accent, text_color)
    else:
        add_bullets(slide, Inches(0.85), Inches(2.0), Inches(10.9), Inches(3.9), body, body_font, 19, text_color)

    notes = slide_spec.get("notes")
    evidence = list(slide_spec.get("evidence", []))
    if notes:
        evidence.append({"source": "notes", "locator": "", "claim": notes})
    add_footer(slide, evidence, body_font, RGBColor(90, 90, 90))


def render(spec: dict, output: Path, template_path: Path | None) -> None:
    prs = Presentation(str(template_path)) if template_path else Presentation()
    if template_path:
        remove_template_slides(prs)
    else:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

    for index, slide_spec in enumerate(spec.get("slides", []), start=1):
        layout = choose_layout(prs, slide_spec) if template_path else prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        populate_slide(slide, slide_spec, spec, index)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec")
    parser.add_argument("--out", required=True)
    parser.add_argument("--template")
    args = parser.parse_args()

    spec = json.loads(Path(args.spec).read_text(encoding="utf-8-sig"))
    if not spec.get("slides"):
        raise SystemExit("Deck spec has no slides.")
    template_path = resolve_template(spec, args.template)
    render(spec, Path(args.out), template_path)
    template_msg = f" using {template_path}" if template_path else " without template"
    print(f"Wrote {args.out}{template_msg}")


if __name__ == "__main__":
    main()
